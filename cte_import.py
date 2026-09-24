"""
Import the team's CTE workbooks (one per year) into the dashboard database.

Reads every "CTE <year>*.xlsx" in the CTE folder — READ ONLY, nothing in
OneDrive is ever written — and loads three sheets:

  Lead Gen             daily activity per agent       -> cte_activity
  My Business          deal log (all fields)           -> cte_deals
  Financial Statement  monthly income/expense lines   -> cte_financials

Each import replaces that file's rows, so re-running after the spreadsheet
changes is always safe. Sync-conflict copies ("...-DESKTOP-XXXX.xlsx")
are skipped.

Normally the Render cron job (fub_agent_activity_pull.py) calls
import_from_onedrive() every run: it downloads the folder from OneDrive via
Microsoft Graph (needs MS_TENANT_ID / MS_CLIENT_ID / MS_CLIENT_SECRET) and
only re-imports files that changed. By hand:

    python cte_import.py --onedrive       # same as the cron job
    python cte_import.py                  # from the local OneDrive sync folder
    python cte_import.py --dry-run        # read + summarize, write nothing

Settings come from the environment or a local .env file (never committed).
"""

import argparse
import glob
import io
import os
import re
import sys
import warnings
from datetime import date, datetime
from urllib.parse import quote

import openpyxl
from openpyxl.utils import get_column_letter
import psycopg2
import requests
from psycopg2.extras import Json

DEFAULT_FOLDER = os.path.expanduser(r"~\OneDrive - Dream Homes Team\CTE FILES")

# Lead Gen header (normalized) -> cte_activity column
ACTIVITY_COLS = {
    "hours": "hours", "dials": "dials", "contacts": "contacts", "nurtures": "nurtures",
    "listing appts set": "listing_appts_set", "listing appts held": "listing_appts_held",
    "listing contracts signed": "listings_signed",
    "buyer appts set": "buyer_appts_set", "buyer appts held": "buyer_appts_held",
    "buyer rep contracts signed": "buyer_reps_signed",
    "written offers": "written_offers", "# of showings": "showings", "open houses held": "open_houses",
}
TEXT_COLS = {"lead gen source": "lead_gen_source", "notes": "notes"}

# My Business columns are in the same place every year (headers are
# reworded now and then), so they're read by column letter.
DEAL_COLS = {
    "W": ("deal_type", "text"), "X": ("status_raw", "text"), "Y": ("address", "text"),
    "Z": ("clients", "text"), "AA": ("source", "text"), "AB": ("signed_date", "date"),
    "AC": ("list_date", "date"), "AD": ("exp_date", "date"), "AE": ("under_contract_date", "date"),
    "AF": ("proj_close_date", "date"), "AG": ("close_date", "date"), "AH": ("list_price", "num"),
    "AI": ("sale_price", "num"), "AK": ("commission_pct", "num"), "AL": ("units", "num"),
    "AM": ("gci", "num"), "AN": ("bonus", "num"), "AO": ("transaction_fee", "num"),
    "AP": ("broker_split", "num"), "AQ": ("royalty_split", "num"), "AR": ("referral", "num"),
    "AS": ("primary_agent", "text"), "AT": ("primary_pct", "num"), "AU": ("primary_gci", "num"),
    "AV": ("secondary_agent", "text"), "AW": ("secondary_pct", "num"), "AX": ("secondary_gci", "num"),
    "AY": ("agent3", "text"), "AZ": ("agent3_pct", "num"), "BA": ("agent3_gci", "num"),
    "BB": ("agent4", "text"), "BC": ("agent4_pct", "num"), "BD": ("agent4_gci", "num"),
    "BE": ("eo_fee", "num"), "BG": ("donation", "num"), "BH": ("other_fee", "num"),
}
DEAL_LAST_COL = 60  # BH; columns past this are CTE's own helper tables


def norm(h):
    return re.sub(r"\s+", " ", str(h or "").replace("*", "")).strip().lower()


def to_num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace("$", "").replace(",", "")
        try:
            return float(s.rstrip("%")) / (100 if s.endswith("%") else 1)
        except ValueError:
            return None
    return None


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def to_text(v):
    if v is None or isinstance(v, bool):
        return None
    s = str(v).strip()
    return None if s == "" or s.startswith("#") else s


def agent_name(v):
    """Agent names are typed by hand; "Corbisiero, Joe" -> "Joe Corbisiero"."""
    s = to_text(v)
    if s and s.count(",") == 1:
        last, first = (x.strip() for x in s.split(","))
        s = f"{first} {last}" if first and last else s
    return re.sub(r"\s+", " ", s) if s else s


def clean_status(s):
    """CTE prefixes statuses to sort them (`Pending, x-Cancelled, z-Expired)."""
    return re.sub(r"^(`|[xyz]-)", "", (s or "").strip()) or None


def jsonable(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float) and v != v:  # NaN
        return None
    return v


def file_year(path):
    m = re.search(r"(20\d\d)", os.path.basename(path))
    return int(m.group(1)) if m else None


def read_activity(wb, year):
    if "Lead Gen" not in wb.sheetnames:
        return [], "no Lead Gen sheet"
    ws = wb["Lead Gen"]
    header = None
    rows = []
    for rownum, r in enumerate(ws.iter_rows(max_col=30, values_only=True), start=1):
        if header is None:
            h = [norm(v) for v in r]
            if "date" in h and "name" in h and "dials" in h:
                header = h
            continue
        rec = dict(zip(header, r))
        d, name = to_date(rec.get("date")), agent_name(rec.get("name"))
        if not d or not name or abs(d.year - year) > 1:
            continue
        row = {"row_num": rownum, "activity_date": d, "agent_name": name}
        extra = {}
        for h, v in rec.items():
            if h in ACTIVITY_COLS:
                row[ACTIVITY_COLS[h]] = to_num(v) or 0
            elif h in TEXT_COLS:
                row[TEXT_COLS[h]] = to_text(v)
            elif h and h not in ("date", "name") and v not in (None, ""):
                extra[h] = jsonable(v)  # custom columns (e.g. "Door Knocks")
        row["extra"] = extra
        rows.append(row)
    return rows, None if header else "old layout (no Date/Name/Dials header), skipped"


def read_deals(wb):
    if "My Business" not in wb.sheetnames:
        return [], "no My Business sheet"
    ws = wb["My Business"]
    header = None
    rows = []
    for rownum, r in enumerate(ws.iter_rows(max_col=DEAL_LAST_COL, values_only=True), start=1):
        if header is None:
            if "Type" in [to_text(v) for v in r] and "Status" in [to_text(v) for v in r]:
                header = [to_text(v) for v in r]
            continue
        cells = {get_column_letter(i + 1): v for i, v in enumerate(r)}
        if to_text(cells.get("W")) not in ("Buyer", "Listing"):
            continue
        row = {"row_num": rownum}
        for col, (name, kind) in DEAL_COLS.items():
            v = cells.get(col)
            row[name] = to_date(v) if kind == "date" else to_num(v) if kind == "num" else to_text(v)
        row["status"] = clean_status(row["status_raw"])
        for k in ("primary_agent", "secondary_agent", "agent3", "agent4"):
            row[k] = agent_name(row[k])
        # Every column W..BH with its header, so nothing in the sheet is lost
        row["raw"] = {f"{get_column_letter(i + 1)} {header[i] or ''}".strip(): jsonable(v)
                      for i, v in enumerate(r) if i >= 22 and v not in (None, "")}
        rows.append(row)
    return rows, None if header else "no Type/Status header found"


def read_financials(wb, year):
    """Financial Statement: every income/expense line with its 12 monthly
    amounts (columns D..O). The table starts at the "Income" row (month-end
    dates across D..O); each line's section is set by CTE's own total rows."""
    if "Financial Statement" not in wb.sheetnames:
        return [], "no Financial Statement sheet"
    ws = wb["Financial Statement"]
    started, section, rows = False, "Income", []
    for rownum, r in enumerate(ws.iter_rows(max_col=17, values_only=True), start=1):
        label = to_text(r[1]) if len(r) > 1 else None
        months = [to_num(v) for v in r[3:15]]
        if not started:
            started = label == "Income" and sum(isinstance(v, (datetime, date)) for v in r[3:15]) >= 10
            continue
        if not label or all(m is None for m in months):
            continue
        low = label.lower()
        if section == "Gross profit" and low in ("operating expenses", "total expenses"):
            section = "Expenses"
        rows.append({"row_num": rownum, "section": section, "label": label,
                     "amounts": [m or 0 for m in months]})
        if low == "total income":
            section = "Income summary"
        elif low == "cost of sales" and section == "Income summary":
            section = "Cost of Sales"
        elif low == "total cost of sales":
            section = "Gross profit"
        elif low == "total expenses":
            section = "Summary"
    return rows, None if started else "no monthly table found"


def ensure_tables(cur):
    num_cols = ", ".join(f"{c} NUMERIC DEFAULT 0" for c in ACTIVITY_COLS.values())
    cur.execute(f"""
        CREATE TABLE IF NOT EXISTS cte_activity (
            source_file TEXT NOT NULL,
            file_year INT NOT NULL,
            row_num INT NOT NULL,
            activity_date DATE NOT NULL,
            agent_name TEXT NOT NULL,
            {num_cols},
            lead_gen_source TEXT,
            notes TEXT,
            extra JSONB,
            PRIMARY KEY (source_file, row_num)
        )
    """)
    deal_cols = ", ".join(
        f"{name} {'DATE' if kind == 'date' else 'NUMERIC' if kind == 'num' else 'TEXT'}"
        for name, kind in DEAL_COLS.values())
    cur.execute(f"""
        CREATE TABLE IF NOT EXISTS cte_deals (
            source_file TEXT NOT NULL,
            file_year INT NOT NULL,
            row_num INT NOT NULL,
            status TEXT,
            {deal_cols},
            raw JSONB,
            PRIMARY KEY (source_file, row_num)
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS cte_financials (
            source_file TEXT NOT NULL,
            file_year INT NOT NULL,
            row_num INT NOT NULL,
            section TEXT,
            label TEXT NOT NULL,
            month INT NOT NULL,
            amount NUMERIC NOT NULL,
            PRIMARY KEY (source_file, row_num, month)
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS cte_import_log (
            source_file TEXT PRIMARY KEY,
            file_year INT,
            file_modified TIMESTAMPTZ,
            imported_at TIMESTAMPTZ DEFAULT NOW(),
            activity_rows INT,
            deal_rows INT
        )
    """)
    cur.execute("ALTER TABLE cte_import_log ADD COLUMN IF NOT EXISTS etag TEXT")


def insert_rows(cur, table, fixed, rows):
    for row in rows:
        data = {**fixed, **row}
        for k in ("extra", "raw"):
            if k in data:
                data[k] = Json(data[k])
        cols = ", ".join(data)
        cur.execute(f"INSERT INTO {table} ({cols}) VALUES ({', '.join(['%s'] * len(data))})",
                    list(data.values()))



def is_cte_file(name):
    """CTE yearly workbooks, minus sync-conflict copies and Excel lock files."""
    return (name.lower().endswith(".xlsx") and name.upper().startswith("CTE ")
            and "-DESKTOP-" not in name.upper() and not name.startswith("~$"))


def import_workbook(cur, name, source, modified=None, etag=None, dry_run=False):
    """Read one CTE workbook (a path or a file-like object) and replace that
    file's rows in the database. Returns a one-line summary."""
    year = file_year(name)
    if not year:
        return f"{name}: no year in the file name, skipped"
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    activity, a_note = read_activity(wb, year)
    deals, d_note = read_deals(wb)
    fin, f_note = read_financials(wb, year)
    wb.close()
    income = next((sum(r["amounts"]) for r in fin if r["label"].lower() == "total income"), 0)
    summary = (f"{name}: {len(activity)} activity rows{f' ({a_note})' if a_note else ''}, "
               f"{len(deals)} deals{f' ({d_note})' if d_note else ''}, "
               f"{len(fin)} financial lines{f' ({f_note})' if f_note else ''}, total income ${income:,.0f}")
    if dry_run:
        return summary
    for table in ("cte_activity", "cte_deals", "cte_financials"):
        cur.execute(f"DELETE FROM {table} WHERE source_file = %s", (name,))
    insert_rows(cur, "cte_activity", {"source_file": name, "file_year": year}, activity)
    insert_rows(cur, "cte_deals", {"source_file": name, "file_year": year}, deals)
    insert_rows(cur, "cte_financials", {"source_file": name, "file_year": year},
                [{"row_num": r["row_num"], "section": r["section"], "label": r["label"],
                  "month": m + 1, "amount": amt} for r in fin for m, amt in enumerate(r["amounts"])])
    cur.execute("""
        INSERT INTO cte_import_log (source_file, file_year, file_modified, etag, imported_at, activity_rows, deal_rows)
        VALUES (%s, %s, %s, %s, NOW(), %s, %s)
        ON CONFLICT (source_file) DO UPDATE SET file_year = EXCLUDED.file_year,
            file_modified = EXCLUDED.file_modified, etag = EXCLUDED.etag, imported_at = NOW(),
            activity_rows = EXCLUDED.activity_rows, deal_rows = EXCLUDED.deal_rows
    """, (name, year, modified, etag, len(activity), len(deals)))
    return summary


# ---------------------------------------------------------------- OneDrive (Microsoft Graph)
# Read-only: the app registration only has Files.Read.All, and this code only
# lists the folder and downloads files.

GRAPH = "https://graph.microsoft.com/v1.0"
DEFAULT_ONEDRIVE_USER = "joecorbisiero@dreamhomesteam.onmicrosoft.com"
DEFAULT_ONEDRIVE_FOLDER = "CTE FILES"


def graph_configured():
    return all(os.environ.get(k) for k in ("MS_TENANT_ID", "MS_CLIENT_ID", "MS_CLIENT_SECRET"))


def graph_session():
    resp = requests.post(
        f"https://login.microsoftonline.com/{os.environ['MS_TENANT_ID']}/oauth2/v2.0/token",
        data={"grant_type": "client_credentials", "client_id": os.environ["MS_CLIENT_ID"],
              "client_secret": os.environ["MS_CLIENT_SECRET"], "scope": "https://graph.microsoft.com/.default"},
        timeout=30)
    if resp.status_code != 200:
        raise RuntimeError(f"Microsoft sign-in failed ({resp.status_code}): {resp.text[:300]}")
    s = requests.Session()
    s.headers["Authorization"] = f"Bearer {resp.json()['access_token']}"
    return s


def onedrive_files(session, user, folder):
    """[{name, id, modified, etag}] for the CTE workbooks in the user's OneDrive folder."""
    url = f"{GRAPH}/users/{quote(user)}/drive/root:/{quote(folder)}:/children?$top=200"
    files = []
    while url:
        resp = session.get(url, timeout=60)
        if resp.status_code == 404:
            raise RuntimeError(f'Folder "{folder}" not found in {user}\'s OneDrive')
        resp.raise_for_status()
        data = resp.json()
        for item in data.get("value", []):
            if "file" in item and is_cte_file(item["name"]):
                files.append({"name": item["name"], "id": item["id"], "etag": item.get("eTag"),
                              "modified": parse_graph_time(item.get("lastModifiedDateTime"))})
        url = data.get("@odata.nextLink")
    return sorted(files, key=lambda f: f["name"])


def parse_graph_time(s):
    return datetime.fromisoformat(s.replace("Z", "+00:00")) if s else None


def import_from_onedrive(database_url, force=False, log=print):
    """Download changed CTE workbooks from OneDrive and import them.
    Unchanged files (same eTag as the last import) are skipped."""
    user = os.environ.get("MS_ONEDRIVE_USER") or DEFAULT_ONEDRIVE_USER
    folder = os.environ.get("CTE_ONEDRIVE_FOLDER") or DEFAULT_ONEDRIVE_FOLDER
    session = graph_session()
    files = onedrive_files(session, user, folder)
    log(f"CTE: {len(files)} workbooks in OneDrive/{folder}")

    conn = psycopg2.connect(database_url)
    cur = conn.cursor()
    ensure_tables(cur)
    conn.commit()
    cur.execute("SELECT source_file, etag FROM cte_import_log")
    seen = dict(cur.fetchall())
    for f in files:
        if not force and f["etag"] and seen.get(f["name"]) == f["etag"]:
            continue
        resp = session.get(f"{GRAPH}/users/{quote(user)}/drive/items/{f['id']}/content", timeout=300)
        resp.raise_for_status()
        log("CTE: " + import_workbook(cur, f["name"], io.BytesIO(resp.content), f["modified"], f["etag"]))
        conn.commit()  # one file at a time, so a bad file doesn't undo the others
    unchanged = sum(1 for f in files if f["etag"] and seen.get(f["name"]) == f["etag"])
    if unchanged and not force:
        log(f"CTE: {unchanged} unchanged since last import, skipped")
    cur.close()
    conn.close()


# ---------------------------------------------------------------- command line

def load_env_file():
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env")
    if os.path.exists(path):
        for line in open(path, encoding="utf-8"):
            if "=" in line and not line.lstrip().startswith("#"):
                k, v = line.strip().split("=", 1)
                os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("folder", nargs="?", default=DEFAULT_FOLDER)
    parser.add_argument("--dry-run", action="store_true", help="read and summarize only")
    parser.add_argument("--onedrive", action="store_true", help="download from OneDrive via Microsoft Graph")
    parser.add_argument("--force", action="store_true", help="with --onedrive: re-import unchanged files too")
    args = parser.parse_args()
    warnings.filterwarnings("ignore")  # openpyxl chatter about data validation / styles
    load_env_file()

    if args.onedrive:
        if not graph_configured() or not os.environ.get("DATABASE_URL"):
            sys.exit("Set MS_TENANT_ID, MS_CLIENT_ID, MS_CLIENT_SECRET and DATABASE_URL first.")
        import_from_onedrive(os.environ["DATABASE_URL"], force=args.force)
        return

    paths = [p for p in sorted(glob.glob(os.path.join(args.folder, "*.xlsx"))) if is_cte_file(os.path.basename(p))]
    if not paths:
        sys.exit(f"No CTE files found in {args.folder}")
    conn = cur = None
    if not args.dry_run:
        if not os.environ.get("DATABASE_URL"):
            sys.exit("Set DATABASE_URL (or put it in a .env file next to this script).")
        conn = psycopg2.connect(os.environ["DATABASE_URL"])
        cur = conn.cursor()
        ensure_tables(cur)
        conn.commit()
    for path in paths:
        modified = datetime.fromtimestamp(os.path.getmtime(path)).astimezone()
        print(import_workbook(cur, os.path.basename(path), path, modified, None, args.dry_run))
        if conn:
            conn.commit()
    if conn:
        conn.close()
    print("Dry run, nothing written." if args.dry_run else "Done.")


if __name__ == "__main__":
    main()
